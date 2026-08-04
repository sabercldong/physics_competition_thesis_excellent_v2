#!/usr/bin/env python3
"""Synchronize the defense deck with the author-authorized analysis table."""

from __future__ import annotations

import os
import re
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor


ROOT = Path(__file__).resolve().parents[2]
PPTX = ROOT / "10_outputs/答辩PPT/全国中学生物理竞赛复赛理论试题的科学思维要求及教学设计启示_答辩PPT.pptx"
FONT_NAME = "Arial Unicode MS"


TEXT_REPLACEMENTS = {
    "逐题人工复核": "独立编码复核",
    "差异记录更正并形成正式版本": "双人独立编码\v146条分歧原文复核",
    "每条记录的核心字段": "独立编码一致性",
    "分析单元 · 维度 · 是否出现 · 等级 · 题干引文 · 解答引文 · 定位信息 · 排除理由":
        "评分者F—G：存在性一致89.67%，Cohen's κ=0.790；五分类加权κ=0.752",
    "91.4%": "86.9%",
    "16.3%": "11.3%",
    "L3：116 / 202": "L3：96 / 192",
    "L3：139 / 221": "L3：126 / 221",
    "76.0%": "77.8%",
    "168个分析单元，构成绝对主流": "172个分析单元，构成绝对主流",
    "14.5%": "8.1%",
    "32个分析单元，加入显性证据链": "18个分析单元，加入显性证据链",
    "7.7%": "10.4%",
    "17个分析单元，模型要素不突出": "23个分析单元，模型要素不突出",
    "四维同时出现：2个单元（0.9%）": "四维同时出现：1个单元（0.5%）",
    "83.9%—96.2%": "78.1%—92.3%",
    "0—31.3%": "0—21.9%",
    "模型91.40% · 推理100% · 论证16.29% · 创新0.91%":
        "模型86.88% · 推理100% · 论证11.31% · 创新0.90%",
    "模型86.88% · 推理100% · 论证11.31% · 创新0.91%":
        "模型86.88% · 推理100% · 论证11.31% · 创新0.90%",
    "模型91.53% · 推理100% · 论证16.87% · 创新0.96%":
        "模型87.11% · 推理100% · 论证11.96% · 创新0.96%",
    "模型92.81% · 推理100% · 论证15.74% · 创新0.97%":
        "模型89.29% · 推理100% · 论证11.40% · 创新0.97%",
    "1.40": "2.41",
    "推理覆盖全部单元，模型建构超过九成，两者共同构成主体结构。":
        "推理覆盖全部单元，模型建构接近九成，两者共同构成主体结构。",
    "形成维度—等级—双端证据的编码路径":
        "双人独立编码与分歧复核",
    "识别系统、变量、边界与理想化条件，选择、建立或调整物理模型。":
        "识别系统、变量、边界与理想化条件，\v选择、建立或调整物理模型。",
    "围绕主张组织证据与理由，进行比较、检验、边界分析或辩护。":
        "围绕主张组织证据与理由，\v进行比较、检验、边界分析或辩护。",
    "评价假设局限、比较替代方案、修正模型或提出开放性设计。":
        "评价假设局限、比较替代方案，\v修正模型或提出开放性设计。",
    "一行对应一个可独立判断的任务": "一行对应一个\v可独立判断的任务",
    "出现性、等级与相邻等级排除": "出现性、等级与\v相邻等级排除",
    "题干触发语 + 解答关键步骤": "题干触发语\v解答关键步骤",
    "双人独立编码 + 146条分歧原文复核": "双人独立编码\v146条分歧原文复核",
    "总体分布、敏感性与典型机制": "总体分布、敏感性\v与典型机制",
    "复杂条件和多条物理关系被整合为连续推理链。":
        "复杂条件和多条物理关系\v被整合为连续推理链。",
    "试题要求很高，但主要仍在给定情境内整合，不普遍要求跨情境重构。":
        "试题要求很高，但主要仍在\v给定情境内整合，不普遍要求\v跨情境重构。",
    "高阶教学设计不能只复制试题结构，需要有意识地补入论证、比较与方案修正。":
        "高阶教学设计不能只复制试题结构，\v需要有意识地补入论证、比较与方案修正。",
    "质量为 M 的天体附近有一束星光，以速度 v₀、瞄准距离 b 掠过。要求在弱偏折近似下建立运动关系并求偏转角。":
        "质量为 M 的天体附近有一束星光，以速度 v₀、瞄准距离 b 掠过。\v要求在弱偏折近似下建立运动关系并求偏转角。",
    "圆柱在可水平移动的斜面上纯滚动。系统同时包含圆柱的平动、转动、斜面的平动以及接触处无滑动约束。":
        "圆柱在可水平移动的斜面上纯滚动。\v系统同时包含圆柱的平动、转动、\v斜面的平动以及接触处无滑动约束。",
    "μ₀ = [2(M+m)/(7M+2m)] tanθ": "μ₀ = [2(M+m)/(7M+2m)] tan θ",
    "改用晶格中的原子核消除显著反冲，并通过多普勒速度扫描补偿微小频移，提出可操作的测量方案。":
        "改用晶格中的原子核消除显著反冲，\v并通过多普勒速度扫描补偿微小频移，\v提出可操作的测量方案。",
    "判断能否共振": "判断共振条件",
    "画出能量区间": "绘制能量区间",
    "主张—证据—理由": "组织证据链",
    "自由核→晶格核": "自由核改晶格核",
    "提出速度扫描方案": "设计速度扫描",
    "精度与可操作性": "比较精度与操作",
    "计算并提交退出条": "修订并总结",
    "它们依赖题目是否显性要求比较、检验、修正模型或设计方案。":
        "它们依赖题目是否显性要求比较、检验、\v修正模型或设计方案。",
    "通过关系图、证据链、模型比较和方案修订，让思维过程成为可观察产出。":
        "通过关系图、证据链、模型比较和方案修订，\v让思维过程成为可观察产出。",
    "形成双人独立编码与分歧证据复核路径": "双人独立编码与分歧复核",
    "建立连续八届、221单元的正式数据集": "连续八届、221单元数据集",
    "用典型题揭示高认知要求的形成机制": "典型题解释高认知要求机制",
    "给出任务链、课堂产出与评价量规": "任务链、课堂产出与评价量规",
    "本研究聚焦试题任务要求；教学设计仍需在真实课堂中继续检验与修订。":
        "本研究聚焦试题任务要求；教学设计仍需在\v真实课堂中继续检验与修订。",
}


NOTE_REPLACEMENTS = {
    "正式版本纳入逐题复核后的更正。":
        "两名匿名评分者对75个单元独立编码，三方分歧再回到完整题干与解答复核。",
    "所有分析单元都要求科学推理，九成以上包含模型建构；科学论证只在约六分之一的单元出现，质疑创新极少。":
        "所有分析单元都要求科学推理，86.88%包含模型建构；科学论证出现于11.31%的单元，质疑创新极少。",
    "约占四分之三": "占77.83%",
    "最大绝对变化为1.40个百分点": "最大绝对变化为2.41个百分点",
    "/07_analysis/results/v3_final/总体维度出现分布.csv":
        "/07_analysis/results/作者授权采用分析版/描述统计/总体维度出现分布.csv",
    "/07_analysis/results/v3_final/各维度等级分布.csv":
        "/07_analysis/results/作者授权采用分析版/描述统计/各维度等级分布.csv",
    "/07_analysis/results/v3_final/维度组合分布.csv":
        "/07_analysis/results/作者授权采用分析版/描述统计/维度组合分布.csv",
    "/07_analysis/results/v3_final/分届次维度出现分布.csv":
        "/07_analysis/results/作者授权采用分析版/描述统计/分届次维度出现分布.csv",
    "/07_analysis/results/v3_final/敏感性分析/三种分析单位权重下四维出现比例.csv":
        "/07_analysis/results/作者授权采用分析版/敏感性分析/三种分析单位权重下四维出现比例.csv",
    "/07_analysis/results/v3_final/":
        "/07_analysis/results/作者授权采用分析版/",
    "/08_manuscript/chapters/chapter1_introduction.md": "/08_manuscript/01_introduction.md",
    "/08_manuscript/chapters/chapter2_literature_review.md": "/08_manuscript/02_theory.md",
    "/08_manuscript/chapters/chapter3_research_design.md": "/08_manuscript/03_methods.md",
    "/08_manuscript/chapters/chapter5_case_analysis.md": "/08_manuscript/05_cases_and_diagnosis.md",
    "/08_manuscript/chapters/chapter6_teaching_implications.md": "/08_manuscript/05_cases_and_diagnosis.md",
    "/08_manuscript/chapters/chapter7_conclusion.md": "/08_manuscript/06_conclusion.md",
    "/05_coding/manual_v3.md": "/05_coding/科学思维编码手册_第三版.md",
    "/06_coded_data/v3_final/": "/07_analysis/data/作者授权采用分析版/",
}


def replace_shape_text(prs: Presentation) -> None:
    def set_single_paragraph_text(shape, value: str) -> None:
        paragraph = shape.text_frame.paragraphs[0]
        runs = paragraph.runs
        if not runs:
            raise RuntimeError(f"cannot safely replace unformatted text: {shape.text}")
        parts = value.split("\v")
        runs[0].text = parts[0]
        for run in runs[1:]:
            run.text = ""
        for part in parts[1:]:
            paragraph.add_line_break()
            paragraph.add_run().text = part

    # Earlier updater versions escaped soft breaks as visible `_x000B_` text.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame or "_x000B_" not in shape.text:
                continue
            if len(shape.text_frame.paragraphs) != 1:
                raise RuntimeError(f"cannot safely repair line breaks: {shape.text}")
            set_single_paragraph_text(shape, shape.text.replace("_x000B_", "\v"))

    pending = dict(TEXT_REPLACEMENTS)
    current_texts = {
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    }
    for old, new in list(pending.items()):
        if old not in current_texts and new in current_texts:
            pending.pop(old)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            old = shape.text
            if old not in pending:
                continue
            new = pending.pop(old)
            paragraphs = shape.text_frame.paragraphs
            if len(paragraphs) != 1 or not paragraphs[0].runs:
                raise RuntimeError(f"cannot safely replace formatted text: {old}")
            set_single_paragraph_text(shape, new)
    remaining_texts = {
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    }
    pending = {old: new for old, new in pending.items() if new not in remaining_texts}
    if pending:
        raise RuntimeError(f"deck text not found: {sorted(pending)}")


def apply_fonts(prs: Presentation) -> None:
    """Use an explicit CJK font so PowerPoint and LibreOffice render identically."""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = FONT_NAME
            if not shape.has_chart:
                continue
            chart = shape.chart
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name = FONT_NAME
            if chart.has_legend:
                chart.legend.font.name = FONT_NAME
            for series in chart.series:
                try:
                    series.data_labels.font.name = FONT_NAME
                except ValueError:
                    pass


def normalize_package_fonts(path: Path) -> None:
    """Set latin/east-Asian/complex-script typefaces in every presentation XML part."""
    font = FONT_NAME.encode("utf-8")
    pattern = re.compile(rb"(<a:(?:latin|ea|cs)\b[^>]*\btypeface=\")[^\"]*(\")")
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".pptx", delete=False) as tmp:
        temp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("ppt/") and info.filename.endswith(".xml"):
                    payload = pattern.sub(lambda match: match.group(1) + font + match.group(2), payload)
                target.writestr(info, payload)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def replace_notes(prs: Presentation) -> None:
    for slide in prs.slides:
        for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
            for run in paragraph.runs:
                text = run.text
                for old, new in NOTE_REPLACEMENTS.items():
                    text = text.replace(old, new)
                run.text = text


def replace_chart_data(prs: Presentation) -> None:
    chart_specs = {
        7: (
            ["模型建构", "科学推理", "科学论证", "质疑创新"],
            [
                ("出现", [86.88, 100.00, 11.31, 0.90]),
                ("未出现", [13.12, 0.00, 88.69, 99.10]),
            ],
        ),
        8: (
            ["模型建构", "科学推理", "科学论证", "质疑创新"],
            [
                ("L1", [14.58, 12.22, 48.00, 0.00]),
                ("L2", [35.42, 30.77, 44.00, 100.00]),
                ("L3", [50.00, 57.01, 8.00, 0.00]),
            ],
        ),
        10: (
            ["35届", "36届", "37届", "38届", "39届", "40届", "41届", "42届"],
            [
                ("模型建构", [92.31, 90.91, 90.91, 92.00, 81.48, 88.89, 78.13, 83.87]),
                ("科学论证", [7.69, 0.00, 13.64, 8.00, 11.11, 13.89, 21.88, 9.68]),
            ],
        ),
    }
    for slide_no, (categories, series) in chart_specs.items():
        charts = [shape.chart for shape in prs.slides[slide_no - 1].shapes if shape.has_chart]
        if len(charts) != 1:
            raise RuntimeError(f"slide {slide_no} expected one chart, found {len(charts)}")
        data = ChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        charts[0].replace_data(data)
        chart = charts[0]
        chart.category_axis.tick_labels.font.name = FONT_NAME
        chart.value_axis.tick_labels.font.name = FONT_NAME
        if chart.has_legend:
            chart.legend.font.name = FONT_NAME
        for chart_series in chart.series:
            try:
                chart_series.data_labels.font.name = FONT_NAME
            except ValueError:
                pass
        if slide_no == 7:
            chart.series[0].data_labels.font.color.rgb = RGBColor(255, 255, 255)
            chart.series[1].data_labels.font.color.rgb = RGBColor(62, 62, 62)
        if slide_no == 8:
            chart.series[0].data_labels.font.color.rgb = RGBColor(62, 62, 62)
            chart.series[1].data_labels.font.color.rgb = RGBColor(62, 62, 62)
            chart.series[2].data_labels.font.color.rgb = RGBColor(255, 255, 255)
        if slide_no == 10:
            for chart_series in chart.series:
                chart_series.smooth = False


def main() -> None:
    prs = Presentation(PPTX)
    replace_shape_text(prs)
    replace_notes(prs)
    replace_chart_data(prs)
    apply_fonts(prs)
    output = PPTX.with_name(PPTX.stem + "_更新中.pptx")
    prs.save(output)
    os.replace(output, PPTX)
    normalize_package_fonts(PPTX)
    print(PPTX)


if __name__ == "__main__":
    main()
