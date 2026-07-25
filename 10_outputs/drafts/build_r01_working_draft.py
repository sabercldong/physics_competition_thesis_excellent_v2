#!/usr/bin/env python3
"""Create a non-submission Word working draft from the current Markdown sources."""

from __future__ import annotations

import hashlib
import json
import re
import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT_DIR = ROOT / "08_manuscript"
OUTPUT_PATH = ROOT / "10_outputs/论文初稿_阅读版/中学生物理竞赛复赛试题的科学思维要求与教学转化研究_论文初稿.docx"
MANIFEST_PATH = ROOT / "10_outputs/论文初稿_阅读版/论文初稿_输入清单.json"
WORKING_TITLE = "中学生物理竞赛复赛试题的科学思维要求与教学转化研究"
WORKING_TITLE_STATUS = "内部工作题目，待作者/导师确认"
CHAPTERS = [
    MANUSCRIPT_DIR / "01_introduction.md",
    MANUSCRIPT_DIR / "02_theory.md",
    MANUSCRIPT_DIR / "03_methods.md",
    MANUSCRIPT_DIR / "04_results.md",
    MANUSCRIPT_DIR / "05_cases_and_diagnosis.md",
    MANUSCRIPT_DIR / "06_conclusion.md",
]
# Songti SC is resolvable by the local Writer renderer and keeps CJK glyphs intact in PDF QA.
FONT_CN = "Songti SC"
FONT_HEADING = "Songti SC"
FONT_EN = "Times New Roman"
REFERENCE_CANDIDATE_PATH = MANUSCRIPT_DIR / "07_candidate_references.md"
CHINESE_ABSTRACT_PATH = MANUSCRIPT_DIR / "00_abstract_content_draft.md"
ENGLISH_ABSTRACT_PATH = MANUSCRIPT_DIR / "00_abstract_en_content_draft.md"
REFERENCE_MAPPING_PATH = ROOT / "09_quality/r01_candidate_references_mapping_v1.csv"
SOURCE_ID = re.compile(r"\b(?:(?:CS|ST|METH|PCOMP)-\d{3}|SOLO-(?:(?:APP|PHY)-\d{3}|\d{3}))\b")
CITATION_PLACEHOLDER = re.compile(r"\[引文占位：([^\]]+)\]")
DISPLAY_REPLACEMENTS = {
    "03_corpus/item_index.csv": "复赛试题小问级索引表",
    "03_corpus/manual_review_checklist.csv": "人工校对清单",
    "07_analysis/data/final_coding_r01_correction_v1.csv": "当前更正编码表",
    "07_analysis/data/final_coding_evidence_r01_v1.csv": "编码证据定位表",
    "07_analysis/data/final_coding.csv": "历史基线编码表",
    "07_analysis/data/internal_coding_v1.csv": "作者预审编码表",
    "07_analysis/data/internal_analysis_candidate_v0.csv": "作者预审分析候选表",
    "final_dimension_presence.csv": "四维度出现描述统计表",
    "final_level_distribution.csv": "等级分布描述统计表",
    "final_dimension_cooccurrence.csv": "维度共现描述统计表",
    "final_year_dimension_presence.csv": "届次维度出现描述表",
    "final_knowledge_dimension_presence.csv": "知识模块维度出现描述表",
    "final_representation_component_dimension_presence.csv": "表征要素分层描述表",
    "final_math_tools_component_dimension_presence.csv": "数学工具要素分层描述表",
    "RATER-A/B": "两名独立编码者",
    "ADJ-01": "首次裁决",
    "ADJ-02": "更正裁决",
    "R01": "研究者复核",
    "参考文献（候选工作稿）": "参考文献",
    "工作稿": "初稿",
}
DISPLAY_SKIP_PHRASES = (
    "当前版本为核验式草稿",
    "本章当前为证据边界版草稿",
    "下一步需要继续完成",
    "当前草稿后续章节将按",
)


def load_candidate_reference_numbers() -> dict[str, int]:
    """Use the reproducible candidate-reference order without changing Markdown evidence IDs."""
    with REFERENCE_MAPPING_PATH.open(encoding="utf-8-sig", newline="") as handle:
        return {
            row["source_id"]: int(row["reference_number"])
            for row in csv.DictReader(handle)
        }


CANDIDATE_REFERENCE_NUMBERS = load_candidate_reference_numbers()


def set_run_font(run, name: str, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), name)
    # LibreOffice can otherwise choose the Latin fallback for CJK runs.
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width: Cm) -> None:
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    field_char = OxmlElement("w:fldChar")
    field_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    run._r.append(field_char)
    run._r.append(instr_text)
    run._r.append(end_char)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(25)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[style_name]
        style.font.name = FONT_HEADING
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(25)
        style.paragraph_format.space_before = Pt(12 if style_name != "Heading 3" else 6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    for name, base in (("Figure Caption", 10.5), ("Table Note", 10.5)):
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = FONT_CN
        style.font.size = Pt(base)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(18)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("华东师范大学硕士专业学位论文")
    set_run_font(header_run, FONT_CN, 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("第 ")
    set_run_font(footer_run, FONT_CN, 9)
    add_page_field(footer)
    tail = footer.add_run(" 页")
    set_run_font(tail, FONT_CN, 9)


def add_cover(document: Document) -> None:
    for _ in range(5):
        document.add_paragraph()
    school = document.add_paragraph()
    school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = school.add_run("华东师范大学")
    set_run_font(run, FONT_HEADING, 26, bold=True)
    degree = document.add_paragraph()
    degree.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = degree.add_run("硕士专业学位论文")
    set_run_font(run, FONT_HEADING, 22, bold=True)
    for _ in range(5):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(WORKING_TITLE)
    set_run_font(run, FONT_HEADING, 18, bold=True)
    for _ in range(6):
        document.add_paragraph()
    metadata = [
        "学位类别：非全日制硕士专业学位",
        "学位类型：教育硕士专业学位（非全日制）",
        "论文初稿（阅读版）",
        "生成日期：2026-07-25",
    ]
    for item in metadata:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(item)
        set_run_font(run, FONT_CN, 12)
    document.add_page_break()


def strip_markdown(text: str) -> str:
    for raw_name, display_name in DISPLAY_REPLACEMENTS.items():
        text = text.replace(raw_name, display_name)
    text = text.replace("R01", "作者预审")
    text = re.sub(r"!\[[^]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def add_inline_text(paragraph, text: str, font_name: str, size: float, bold: bool = False) -> None:
    """Render candidate citations as Word superscript runs, not ordinary body text."""
    cursor = 0
    for match in CITATION_PLACEHOLDER.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(strip_markdown(text[cursor:match.start()]))
            set_run_font(run, font_name, size, bold=bold)
        source_ids = SOURCE_ID.findall(match.group(1))
        numbers = [CANDIDATE_REFERENCE_NUMBERS[source_id] for source_id in source_ids if source_id in CANDIDATE_REFERENCE_NUMBERS]
        unresolved = [source_id for source_id in source_ids if source_id not in CANDIDATE_REFERENCE_NUMBERS]
        if numbers:
            run = paragraph.add_run("[" + ", ".join(str(number) for number in numbers) + "]")
            set_run_font(run, font_name, size, bold=bold)
            run.font.superscript = True
        if unresolved:
            run = paragraph.add_run("[待补核验文献：" + "、".join(unresolved) + "]")
            set_run_font(run, font_name, size, bold=bold)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(strip_markdown(text[cursor:]))
        set_run_font(run, font_name, size, bold=bold)


def add_text_paragraph(document: Document, text: str, style: str | None = None, indent: bool = True) -> None:
    paragraph = document.add_paragraph(style=style)
    if not indent:
        paragraph.paragraph_format.first_line_indent = Cm(0)
    add_inline_text(paragraph, text, FONT_CN, 12)


def add_english_text_paragraph(document: Document, text: str, indent: bool = True) -> None:
    paragraph = document.add_paragraph()
    if not indent:
        paragraph.paragraph_format.first_line_indent = Cm(0)
    add_inline_text(paragraph, text, FONT_EN, 12)


def add_source_numbered_paragraph(document: Document, text: str) -> None:
    """Preserve source numbering without letting Word continue it across chapters."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.74)
    run = paragraph.add_run(strip_markdown(text))
    set_run_font(run, FONT_CN, 12)


def add_contents_entry(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm((level - 1) * 0.74)
    run = paragraph.add_run(strip_markdown(text))
    set_run_font(run, FONT_CN, 12 if level == 1 else 11)


def add_reference_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.74)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(20)
    run = paragraph.add_run(strip_markdown(text))
    set_run_font(run, FONT_CN, 10.5)


def add_markdown_table(document: Document, rows: list[str]) -> None:
    matrix = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows if "---" not in row]
    if not matrix:
        return
    table = document.add_table(rows=len(matrix), cols=len(matrix[0]))
    table.style = "Table Grid"
    table.autofit = False
    total_width = Cm(15.4)
    width = Cm(total_width.cm / len(matrix[0]))
    for row_index, row in enumerate(matrix):
        for column_index, cell_text in enumerate(row):
            cell = table.cell(row_index, column_index)
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(18)
            add_inline_text(paragraph, cell_text, FONT_CN, 9.5, bold=row_index == 0)
            if row_index == 0:
                shade_cell(cell, "D9E2F3")
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def add_image(document: Document, line: str) -> bool:
    match = re.search(r"!\[[^]]*\]\(([^)]+)\)", line)
    if not match:
        return False
    image_path = Path(match.group(1))
    if not image_path.exists():
        return False
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(15.2))
    return True


def add_chapter(document: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if path == REFERENCE_CANDIDATE_PATH and line == "## 使用边界":
            break
        if any(phrase in line for phrase in DISPLAY_SKIP_PHRASES):
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            document.add_heading(strip_markdown(line[level:].strip()), level=level)
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue
        elif line.startswith("!["):
            if not add_image(document, line):
                add_text_paragraph(document, "[图件路径待确认]", style="Figure Caption", indent=False)
        elif re.match(r"^图\d+-\d+", line):
            add_text_paragraph(document, line, style="Figure Caption", indent=False)
        elif re.match(r"^\[\d+\]\s+", line):
            add_reference_paragraph(document, line)
        elif line.startswith(">"):
            index += 1
            continue
        elif re.match(r"^-\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline_text(paragraph, re.sub(r"^-\s+", "", line), FONT_CN, 12)
        elif re.match(r"^\d+\.\s+", line):
            add_source_numbered_paragraph(document, line)
        else:
            add_text_paragraph(document, line)
        index += 1


def main() -> None:
    document = Document()
    configure_document(document)
    add_cover(document)

    abstract_heading = document.add_heading("摘 要", level=1)
    abstract_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in CHINESE_ABSTRACT_PATH.read_text(encoding="utf-8").splitlines()[2:]:
        if line.strip():
            add_text_paragraph(document, line, indent=False)
    document.add_page_break()

    english_abstract_heading = document.add_heading("ABSTRACT", level=1)
    english_abstract_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in english_abstract_heading.runs:
        set_run_font(run, FONT_EN, 16, bold=True)
    for line in ENGLISH_ABSTRACT_PATH.read_text(encoding="utf-8").splitlines()[2:]:
        if line.strip():
            add_english_text_paragraph(document, line, indent=False)
    document.add_page_break()

    contents = document.add_heading("目 录", level=1)
    contents.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for chapter in CHAPTERS:
        for line in chapter.read_text(encoding="utf-8").splitlines():
            match = re.match(r"^(#{1,3})\s+(.+)$", line)
            if not match:
                continue
            level = len(match.group(1))
            title = match.group(2).strip()
            if title == "数据来源":
                continue
            add_contents_entry(document, title, level)
    document.add_page_break()

    for index, chapter in enumerate(CHAPTERS):
        add_chapter(document, chapter)
        if index < len(CHAPTERS) - 1:
            document.add_page_break()

    document.add_page_break()
    add_chapter(document, REFERENCE_CANDIDATE_PATH)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    manifest = {
        "artifact": str(OUTPUT_PATH.relative_to(ROOT)),
        "artifact_scope": "readable thesis draft; not a submission copy",
        "generated_on": "2026-07-25",
        "source_sha256": {
            str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in [
                *CHAPTERS,
                CHINESE_ABSTRACT_PATH,
                ENGLISH_ABSTRACT_PATH,
                REFERENCE_CANDIDATE_PATH,
                REFERENCE_MAPPING_PATH,
                MANUSCRIPT_DIR / "00_frontmatter.md",
                Path(__file__),
            ]
        },
        "figure_inputs": [
            "08_manuscript/figures/figure_1_1_r01_technical_route_v1.png",
            "08_manuscript/figures/figure_2_1_r01_theoretical_framework_v1.png",
            "08_manuscript/figures/figure_4_1_dimension_presence_r01_v1.png",
            "08_manuscript/figures/figure_4_2_level_distribution_r01_v1.png",
        ],
        "exclusions": [
            "personal information",
            "advisor information",
            "student number",
            "signature pages",
            "defense or anonymous-submission versioning",
        ],
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
